import re
import os
import docx
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = 'Times New Roman'
TOTAL_W_DXA = 7937  # 14.0 cm (A4 11906 - 2268 left 4cm - 1701 right 3cm)

def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f"""
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    """)
    tblPr.append(borders)

def add_page_number_to_section(section):
    # Enable different first page
    sectPr = section._sectPr
    titlePg = OxmlElement('w:titlePg')
    sectPr.append(titlePg)

    # Footer for subsequent pages
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    # Add page number run
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    p._p.append(fldSimple)

def plain(s):
    return re.sub(r'[*_`]', '', s)

def col_widths(rows):
    n = len(rows[0])
    def cw(ch):
        if re.match(r'[A-Z0-9#]', ch):
            return 140
        elif re.match(r'[ilftjr.,:;()/\-]', ch):
            return 62
        else:
            return 98

    mins = []
    weights = []
    for c in range(n):
        longest = 0
        tot = 0
        for ri, r in enumerate(rows):
            cell_str = plain(r[c] if c < len(r) else '')
            tot += len(cell_str)
            f = 1.12 if (ri == 0 or (c < len(r) and r[c].startswith('**'))) else 1.0
            for w in cell_str.split():
                w_len = f * sum(cw(ch) for ch in w)
                if w_len > longest:
                    longest = w_len
        mins.append(min(longest * 1.0 + 170, 2500))
        weights.append(max(8, min(tot / max(len(rows), 1), 120)))

    sum_min = sum(mins)
    if sum_min >= TOTAL_W_DXA:
        w = [int(m * TOTAL_W_DXA / sum_min) for m in mins]
    else:
        extra = TOTAL_W_DXA - sum_min
        sw = sum(weights)
        w = [int(m + extra * weights[i] / sw) for i, m in enumerate(mins)]
    
    # Adjust rounding remainder
    rem = TOTAL_W_DXA - sum(w)
    w[-1] += rem
    return w

def runs(paragraph, text, base_font_size=12, default_bold=False, default_italic=False, default_color=None):
    re_tokens = re.compile(r'(\*\*\*([^*]+)\*\*\*|\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`)')
    last = 0
    for m in re_tokens.finditer(text):
        if m.start() > last:
            r = paragraph.add_run(text[last:m.start()])
            r.font.name = FONT_NAME
            r.font.size = Pt(base_font_size)
            if default_bold: r.font.bold = True
            if default_italic: r.font.italic = True
            if default_color: r.font.color.rgb = default_color

        if m.group(2) is not None:
            r = paragraph.add_run(m.group(2))
            r.font.name = FONT_NAME
            r.font.size = Pt(base_font_size)
            r.font.bold = True
            r.font.italic = True
        elif m.group(3) is not None:
            r = paragraph.add_run(m.group(3))
            r.font.name = FONT_NAME
            r.font.size = Pt(base_font_size)
            r.font.bold = True
        elif m.group(4) is not None:
            r = paragraph.add_run(m.group(4))
            r.font.name = FONT_NAME
            r.font.size = Pt(base_font_size)
            r.font.italic = True
        elif m.group(5) is not None:
            r = paragraph.add_run(m.group(5))
            r.font.name = 'Consolas'
            r.font.size = Pt(base_font_size - 1)
        last = m.end()

    if last < len(text):
        r = paragraph.add_run(text[last:])
        r.font.name = FONT_NAME
        r.font.size = Pt(base_font_size)
        if default_bold: r.font.bold = True
        if default_italic: r.font.italic = True
        if default_color: r.font.color.rgb = default_color

def body_p(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=0, space_before=0, line_spacing=1.5, bold=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    runs(p, text, base_font_size=12, default_bold=bold)
    return p

def spacer_p(doc, pt=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pt)
    p.paragraph_format.line_spacing = 1.0

def make_table(doc, rows):
    widths_dxa = col_widths(rows)
    table = doc.add_table(rows=len(rows), cols=len(widths_dxa))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="000000", sz="4", val="single")

    for ri, row in enumerate(rows):
        trow = table.rows[ri]
        trPr = trow._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        if ri == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for ci, cell_text in enumerate(row):
            cell = trow.cells[ci]
            cell.width = Inches(widths_dxa[ci] / 1440.0)
            set_cell_margins(cell, top=40, bottom=40, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            clean_text = cell_text.replace('<br>', '\n').replace('&amp;', '&')
            sub_lines = clean_text.split('\n')
            for sli, sl in enumerate(sub_lines):
                if sli > 0:
                    p = cell.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                runs(p, sl, base_font_size=10.5 if ri > 0 else 10.5, default_bold=(ri == 0))
    return table

def build_srs_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        raw_text = f.read()

    raw_lines = raw_text.split('\n')

    # Place Referensi Acuan Pembelajaran before Lampiran A so body order matches Daftar Isi
    ref_idx = -1
    lamp_idx = -1
    for idx, l in enumerate(raw_lines):
        if l.startswith('## Referensi Acuan') and ref_idx == -1:
            ref_idx = idx
        elif l.startswith('## Lampiran A') and lamp_idx == -1:
            lamp_idx = idx

    if ref_idx > lamp_idx and lamp_idx > 0:
        ref_block = raw_lines[ref_idx:]
        raw_lines = raw_lines[:lamp_idx] + ref_block + ['', '---', ''] + raw_lines[lamp_idx:ref_idx]

    doc = Document()

    # Page Setup (A4, Left 4cm, Right 3cm, Top 4cm, Bottom 3cm)
    sec = doc.sections[0]
    sec.page_width = Inches(21.0 / 2.54)
    sec.page_height = Inches(29.7 / 2.54)
    sec.top_margin = Inches(4.0 / 2.54)
    sec.bottom_margin = Inches(3.0 / 2.54)
    sec.left_margin = Inches(4.0 / 2.54)
    sec.right_margin = Inches(3.0 / 2.54)
    add_page_number_to_section(sec)

    # Styles
    normal = doc.styles['Normal']
    normal.font.name = FONT_NAME
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    normal.paragraph_format.line_spacing = 1.5

    toc_entries = []
    cover = True
    in_table = False
    table_rows = []
    i = 0

    while i < len(raw_lines):
        line = raw_lines[i]
        stripped = line.strip()

        # Check table
        if stripped.startswith('|') and '|' in stripped[1:]:
            in_table = True
            table_rows.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                # Filter separator rows
                cleaned = []
                for r in table_rows:
                    if re.match(r'^\s*\|?\s*[-:\s|]+\s*\|?\s*$', r):
                        continue
                    cells = [c.strip() for c in r.strip().strip('|').split('|')]
                    cleaned.append(cells)
                if cleaned:
                    spacer_p(doc, 4)
                    make_table(doc, cleaned)
                    spacer_p(doc, 4)
                in_table = False
                table_rows = []

        if not stripped:
            i += 1
            continue

        if re.match(r'^---+\s*$', stripped):
            i += 1
            continue

        # Headings
        h_match = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if h_match:
            lvl = len(h_match.group(1))
            text = h_match.group(2).strip()

            if text == 'Daftar Isi':
                cover = False
                doc.add_page_break()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
                runs(p, text, base_font_size=12, default_bold=True)
                i += 1

                # Parse TOC lines
                while i < len(raw_lines) and not re.match(r'^---+\s*$', raw_lines[i].strip()):
                    toc_m = re.match(r'^(\s*)([*-]|\d+\.)\s+(.*)$', raw_lines[i])
                    if toc_m:
                        num = re.search(r'\d+\.', toc_m.group(2))
                        is_sub = (not num and len(toc_m.group(1)) >= 2)
                        t_title = f"{toc_m.group(2)} {toc_m.group(3)}" if num else toc_m.group(3)
                        toc_entries.append((t_title, 1 if is_sub else 0))
                    i += 1

                # Render TOC with tab stops & dot leaders
                for title, indent_level in toc_entries:
                    tp = doc.add_paragraph()
                    tp.paragraph_format.line_spacing = 1.5
                    tp.paragraph_format.space_before = Pt(0)
                    tp.paragraph_format.space_after = Pt(0)
                    tp.paragraph_format.left_indent = Inches(0.5) if indent_level == 1 else Inches(0)
                    
                    # Add right-aligned tab stop with dot leader at TOTAL_W_DXA
                    tp.paragraph_format.tab_stops.add_tab_stop(Inches(TOTAL_W_DXA / 1440.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                    
                    runs(tp, title, base_font_size=12)
                    r_tab = tp.add_run('\t')
                    r_tab.font.name = FONT_NAME
                    r_tab.font.size = Pt(12)
                continue

            if cover:
                # Cover page elements
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_before = Pt(12 if (lvl == 3 and text.startswith('Tim')) else 2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
                runs(p, text, base_font_size=14 if lvl == 1 else 12, default_bold=True)
                i += 1
                continue
            else:
                # Regular chapters and sections
                pb = bool(re.match(r'^(1\. Pendahuluan|Referensi Acuan|Lampiran )', text))
                if pb:
                    doc.add_page_break()
                
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
                runs(p, text, base_font_size=12, default_bold=True)
                i += 1
                continue

        # Code block & HTML details tags (skip from printable DOCX body)
        if stripped.startswith('<details') or stripped.startswith('</details') or stripped.startswith('<summary') or stripped.startswith('</summary'):
            i += 1
            continue

        if stripped.startswith('```'):
            i += 1
            while i < len(raw_lines) and not raw_lines[i].strip().startswith('```'):
                i += 1
            i += 1
            continue

        # Images
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            alt_text, img_file = img_match.groups()
            png_file = re.sub(r'\.svg$', '.png', img_file)
            target_file = png_file if os.path.exists(os.path.join(os.path.dirname(md_path), png_file)) else img_file
            img_path = os.path.join(os.path.dirname(md_path), target_file)
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(6)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.keep_with_next = True
                r_img = p_img.add_run()
                r_img.add_picture(img_path, width=Inches(TOTAL_W_DXA / 1440.0))
            i += 1
            continue

        # Checklists: - [ ] or - [x]
        chk_match = re.match(r'^(\s*)-\s+\[([ xX])\]\s+(.*)$', stripped)
        if chk_match:
            box = '[√]' if chk_match.group(2) in ['x', 'X'] else '[  ]'
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.4)

            r_box = p.add_run(f"{box}  ")
            r_box.font.name = FONT_NAME
            r_box.font.size = Pt(12)
            runs(p, chk_match.group(3), base_font_size=12)
            i += 1
            continue

        # Lists: * or - or 1.
        li_match = re.match(r'^(\s*)([*-]|\d+\.)\s+(.*)$', stripped)
        if li_match:
            indent_spaces = len(li_match.group(1))
            bullet = li_match.group(2)
            content = li_match.group(3)

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            base_ind = 0.3 * (1 + (indent_spaces // 2))
            p.paragraph_format.left_indent = Inches(base_ind)
            p.paragraph_format.first_line_indent = Inches(-0.25)

            marker = f"{bullet} " if re.match(r'\d+\.', bullet) else "•  "
            r_mark = p.add_run(marker)
            r_mark.font.name = FONT_NAME
            r_mark.font.size = Pt(12)
            if re.match(r'\d+\.', bullet):
                r_mark.font.bold = True

            runs(p, content, base_font_size=12)
            i += 1
            continue

        # Captions
        if stripped.startswith('*Gambar') or stripped.startswith('*Prinsip:'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            runs(p, stripped, base_font_size=11, default_italic=True)
            i += 1
            continue

        # Blockquote / Callout
        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.4)
            runs(p, stripped[2:].strip(), base_font_size=11.5, default_italic=True)
            i += 1
            continue

        # Regular Body Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        runs(p, stripped, base_font_size=12)
        i += 1

    if in_table and table_rows:
        cleaned = []
        for r in table_rows:
            if re.match(r'^\s*\|?\s*[-:\s|]+\s*\|?\s*$', r):
                continue
            cells = [c.strip() for c in r.strip().strip('|').split('|')]
            cleaned.append(cells)
        if cleaned:
            spacer_p(doc, 4)
            make_table(doc, cleaned)
            spacer_p(doc, 4)

    doc.save(docx_path)
    print(f"Successfully generated DOCX: {docx_path} ({os.path.getsize(docx_path)} bytes)")

if __name__ == '__main__':
    script_dir = os.path.dirname(os.path.abspath(__file__))
    md = os.path.join(script_dir, "Draft_SRS_Medkreminfo.md")
    out_docx = os.path.join(script_dir, "Draft_SRS_Medkreminfo.docx")
    build_srs_docx(md, out_docx)
    out_rtm = os.path.join(script_dir, "Tugas_2_KelompokXX_DraftSRS.docx")
    build_srs_docx(md, out_rtm)
